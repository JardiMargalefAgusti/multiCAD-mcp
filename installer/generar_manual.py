"""Genera el manual de usuario en Word y PDF a partir de MANUAL_DE_USO.md.

Fuente unica: el .md. Este script lo maqueta con la identidad de APOGEA
(portada, indice, encabezados, tablas con cabecera navy) y exporta a PDF con
Word.

Uso:
    .venv\\Scripts\\python.exe installer\\generar_manual.py

Requisitos: python-docx  (uv pip install python-docx) y Microsoft Word para
el PDF. Sin Word se genera igualmente el .docx.
"""

import datetime
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
ORIGEN = AQUI / "MANUAL_DE_USO.md"
DOCX = AQUI / "MANUAL_DE_USO.docx"
PDF = AQUI / "MANUAL_DE_USO.pdf"

# Paleta corporativa
NAVY = RGBColor(0x0A, 0x16, 0x33)
AZUL = RGBColor(0x1F, 0x7A, 0xE0)
GRIS = RGBColor(0x55, 0x55, 0x55)
CODIGO = RGBColor(0xB1, 0x32, 0x5C)
NAVY_HEX = "0A1633"
GRIS_HEX = "F2F5FA"

VERSION = "0.2.0"


# ========== utilidades de bajo nivel ==========


def _sombrear(celda, color_hex: str) -> None:
    """Rellena una celda con un color solido."""
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), color_hex)
    propiedades.append(sombra)


def _fila_indivisible(fila, es_cabecera: bool = False) -> None:
    """Evita que una fila se parta entre paginas; repite la cabecera."""
    propiedades = fila._tr.get_or_add_trPr()
    no_partir = OxmlElement("w:cantSplit")
    propiedades.append(no_partir)
    if es_cabecera:
        cabecera = OxmlElement("w:tblHeader")
        propiedades.append(cabecera)


def _campo(parrafo, instruccion: str, marcador: str = "") -> None:
    """Inserta un campo de Word (TOC, PAGE...) en el parrafo."""
    ejecucion = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instruccion
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    provisional = OxmlElement("w:t")
    provisional.text = marcador
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for nodo in (inicio, texto, separador, provisional, fin):
        ejecucion._r.append(nodo)


# ========== formato en linea ==========

_TROZOS = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")

# Linea que es integramente un ejemplo de peticion: *"Dibuja un circulo..."*
_EJEMPLO = re.compile(r'^\*".+"\*$')


def escribir_inline(parrafo, texto: str) -> None:
    """Vuelca texto markdown con negritas, cursivas y codigo."""
    for trozo in _TROZOS.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            ejecucion = parrafo.add_run(trozo[1:-1])
            ejecucion.font.name = "Consolas"
            ejecucion.font.size = Pt(9.5)
            ejecucion.font.color.rgb = CODIGO
        elif trozo.startswith("*") and trozo.endswith("*"):
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo)


# ========== estilos del documento ==========


def preparar_estilos(documento: Document) -> None:
    """Tipografia y encabezados con la identidad de APOGEA."""
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for nombre, tamano, color, antes in (
        ("Heading 1", 19, NAVY, 20),
        ("Heading 2", 14, AZUL, 15),
        ("Heading 3", 11.5, NAVY, 11),
    ):
        estilo = documento.styles[nombre]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = color
        estilo.paragraph_format.space_before = Pt(antes)
        estilo.paragraph_format.space_after = Pt(5)
        estilo.paragraph_format.keep_with_next = True

    seccion = documento.sections[0]
    seccion.top_margin = Cm(2.2)
    seccion.bottom_margin = Cm(2.0)
    seccion.left_margin = Cm(2.4)
    seccion.right_margin = Cm(2.4)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ejecucion = pie.add_run("multiCAD-MCP · Manual de uso · ")
    ejecucion.font.size = Pt(8)
    ejecucion.font.color.rgb = GRIS
    _campo(pie, "PAGE", "1")
    for run in pie.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS


def portada(documento: Document, titulo: str, entradilla: list[str]) -> None:
    """Titulo, subtitulo y metadatos."""
    espacio = documento.add_paragraph()
    espacio.paragraph_format.space_after = Pt(90)

    parrafo = documento.add_paragraph()
    ejecucion = parrafo.add_run(titulo)
    ejecucion.font.size = Pt(30)
    ejecucion.font.bold = True
    ejecucion.font.color.rgb = NAVY
    parrafo.paragraph_format.space_after = Pt(4)

    linea = documento.add_paragraph()
    ejecucion = linea.add_run(
        "Control de AutoCAD · ZWCAD · GstarCAD · BricsCAD desde Claude"
    )
    ejecucion.font.size = Pt(13)
    ejecucion.font.color.rgb = AZUL
    linea.paragraph_format.space_after = Pt(22)

    for bloque in entradilla:
        parrafo = documento.add_paragraph()
        escribir_inline(parrafo, bloque)

    fecha = datetime.date.today().strftime("%d/%m/%Y")
    meta = documento.add_paragraph()
    ejecucion = meta.add_run(f"Versión {VERSION} · {fecha} · APOGEA Consulting")
    ejecucion.font.size = Pt(9)
    ejecucion.font.color.rgb = GRIS
    meta.paragraph_format.space_before = Pt(30)

    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Titulo sin estilo Heading: si no, el indice se listaria a si mismo
    rotulo = documento.add_paragraph()
    ejecucion = rotulo.add_run("Contenido")
    ejecucion.font.size = Pt(19)
    ejecucion.font.bold = True
    ejecucion.font.color.rgb = NAVY
    rotulo.paragraph_format.space_after = Pt(10)

    indice = documento.add_paragraph()
    _campo(indice, 'TOC \\o "1-3" \\h \\z \\u', "Actualiza el índice con F9")
    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ========== conversion markdown -> docx ==========


def _es_separador_tabla(linea: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def _celdas(linea: str) -> list[str]:
    return [celda.strip() for celda in linea.strip().strip("|").split("|")]


def anadir_tabla(documento: Document, filas: list[str]) -> None:
    """Convierte un bloque de tabla markdown en tabla de Word."""
    cabecera = _celdas(filas[0])
    datos = [_celdas(fila) for fila in filas[2:]]

    tabla = documento.add_table(rows=1, cols=len(cabecera))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT

    for indice, texto in enumerate(cabecera):
        celda = tabla.rows[0].cells[indice]
        celda.text = ""
        parrafo = celda.paragraphs[0]
        parrafo.paragraph_format.space_after = Pt(2)
        escribir_inline(parrafo, texto)
        for ejecucion in parrafo.runs:
            ejecucion.bold = True
            ejecucion.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            ejecucion.font.size = Pt(10)
        _sombrear(celda, NAVY_HEX)
    _fila_indivisible(tabla.rows[0], es_cabecera=True)

    for numero, fila in enumerate(datos):
        nueva = tabla.add_row()
        _fila_indivisible(nueva)
        celdas = nueva.cells
        for indice, texto in enumerate(fila[: len(cabecera)]):
            celdas[indice].text = ""
            parrafo = celdas[indice].paragraphs[0]
            parrafo.paragraph_format.space_after = Pt(2)
            escribir_inline(parrafo, texto)
            for ejecucion in parrafo.runs:
                ejecucion.font.size = Pt(10)
            if numero % 2 == 1:
                _sombrear(celdas[indice], GRIS_HEX)

    documento.add_paragraph().paragraph_format.space_after = Pt(2)


def anadir_ejemplo(documento: Document, texto: str) -> None:
    """Ejemplo de peticion: una linea propia, en azul y sangrada."""
    parrafo = documento.add_paragraph()
    parrafo.paragraph_format.left_indent = Cm(0.5)
    parrafo.paragraph_format.space_after = Pt(2)
    escribir_inline(parrafo, texto)
    for ejecucion in parrafo.runs:
        ejecucion.italic = True
        ejecucion.font.color.rgb = AZUL


def anadir_cita(documento: Document, texto: str) -> None:
    """Bloque destacado (>) con barra lateral azul."""
    parrafo = documento.add_paragraph()
    parrafo.paragraph_format.left_indent = Cm(0.6)
    parrafo.paragraph_format.space_before = Pt(6)
    parrafo.paragraph_format.space_after = Pt(10)
    escribir_inline(parrafo, texto)
    for ejecucion in parrafo.runs:
        ejecucion.font.color.rgb = AZUL
        ejecucion.italic = True
    bordes = OxmlElement("w:pBdr")
    izquierda = OxmlElement("w:left")
    izquierda.set(qn("w:val"), "single")
    izquierda.set(qn("w:sz"), "18")
    izquierda.set(qn("w:space"), "8")
    izquierda.set(qn("w:color"), "1F7AE0")
    bordes.append(izquierda)
    parrafo._p.get_or_add_pPr().append(bordes)


def convertir(lineas: list[str], documento: Document) -> None:
    """Recorre el markdown y lo va volcando al documento."""
    indice = 0
    while indice < len(lineas):
        linea = lineas[indice].rstrip()

        if not linea.strip() or linea.strip() == "---":
            indice += 1
            continue

        # Tablas
        if linea.startswith("|") and indice + 1 < len(lineas) and _es_separador_tabla(lineas[indice + 1]):
            bloque = []
            while indice < len(lineas) and lineas[indice].strip().startswith("|"):
                bloque.append(lineas[indice])
                indice += 1
            anadir_tabla(documento, bloque)
            continue

        # Encabezados
        encabezado = re.match(r"^(#{1,3})\s+(.*)$", linea)
        if encabezado:
            nivel = len(encabezado.group(1))
            parrafo = documento.add_heading("", level=nivel)
            escribir_inline(parrafo, encabezado.group(2))
            indice += 1
            continue

        # Ejemplos de peticion (cada uno en su linea)
        if _EJEMPLO.match(linea.strip()):
            anadir_ejemplo(documento, linea.strip())
            indice += 1
            continue

        # Citas
        if linea.startswith("> "):
            anadir_cita(documento, linea[2:])
            indice += 1
            continue

        # Listas
        vineta = re.match(r"^[-*]\s+(.*)$", linea)
        numerada = re.match(r"^\d+\.\s+(.*)$", linea)
        if vineta or numerada:
            contenido = (vineta or numerada).group(1)
            # continuaciones indentadas de la misma vineta
            while indice + 1 < len(lineas) and re.match(r"^\s{2,}\S", lineas[indice + 1]):
                indice += 1
                contenido += " " + lineas[indice].strip()
            estilo = "List Number" if numerada else "List Bullet"
            parrafo = documento.add_paragraph(style=estilo)
            parrafo.paragraph_format.space_after = Pt(3)
            escribir_inline(parrafo, contenido)
            indice += 1
            continue

        # Parrafo normal (une lineas hasta el siguiente hueco)
        bloque = [linea]
        while (
            indice + 1 < len(lineas)
            and lineas[indice + 1].strip()
            and not re.match(r"^(#{1,3}\s|[-*]\s|\d+\.\s|\||>\s|---)", lineas[indice + 1])
            and not _EJEMPLO.match(lineas[indice + 1].strip())
        ):
            indice += 1
            bloque.append(lineas[indice].strip())
        escribir_inline(documento.add_paragraph(), " ".join(bloque))
        indice += 1


def generar_docx() -> Path:
    """Construye el .docx completo."""
    lineas = ORIGEN.read_text(encoding="utf-8").splitlines()

    titulo = lineas[0].lstrip("# ").strip()
    # La entradilla es todo lo que hay hasta el primer separador
    corte = next(i for i, l in enumerate(lineas) if l.strip() == "---")

    # La entradilla va hasta el primer separador: un parrafo por bloque, no
    # una linea por linea del markdown
    entradilla: list[str] = []
    bloque: list[str] = []
    for linea in lineas[1:corte]:
        if linea.strip():
            bloque.append(linea.strip())
        elif bloque:
            entradilla.append(" ".join(bloque))
            bloque = []
    if bloque:
        entradilla.append(" ".join(bloque))

    documento = Document()
    preparar_estilos(documento)
    portada(documento, titulo, entradilla)
    convertir(lineas[corte + 1 :], documento)

    documento.core_properties.title = titulo
    documento.core_properties.author = "APOGEA Consulting"
    documento.core_properties.comments = "Manual de uso de multiCAD-MCP"
    documento.save(DOCX)
    return DOCX


def generar_pdf() -> Path | None:
    """Exporta el .docx a PDF con Word (actualizando indice y campos)."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print("pywin32 no disponible: no se genera PDF")
        return None

    pythoncom.CoInitialize()
    word = None
    documento = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        documento = word.Documents.Open(str(DOCX))
        # Indice y numeros de pagina
        for tabla in documento.TablesOfContents:
            tabla.Update()
        documento.Fields.Update()
        documento.ExportAsFixedFormat(
            OutputFileName=str(PDF),
            ExportFormat=17,          # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=0,            # print
            CreateBookmarks=1,        # marcadores desde los encabezados
        )
        documento.Save()
        return PDF
    finally:
        if documento is not None:
            documento.Close(0)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def main() -> int:
    """Genera Word y PDF."""
    if not ORIGEN.exists():
        print(f"No encuentro {ORIGEN}")
        return 1
    ruta_docx = generar_docx()
    print(f"Word: {ruta_docx}  ({ruta_docx.stat().st_size // 1024} KB)")
    ruta_pdf = generar_pdf()
    if ruta_pdf and ruta_pdf.exists():
        print(f"PDF : {ruta_pdf}  ({ruta_pdf.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
